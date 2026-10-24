"""Refine the report: reasons for each design choice and the alternatives, how every metric is calculated,
and what each stage produced.  Original wording is used for the rewritten generic paragraphs.

Input : a report .docx made by scripts/08_make_report.py (default docs/Xtracts_Mini_Project_Report_UPDATED_NEW.docx)
Output: docs/Xtracts_Mini_Project_Report_REFINED.docx   (the input is never modified)

Numbers come from the input report itself and from reports/evaluation/metrics.json (your own evaluation run).
Run it again after a new evaluation: it always starts from the input file, so nothing is duplicated.

Usage:  python scripts/09_refine_report.py [--in FILE] [--out FILE] [--metrics reports/evaluation/metrics.json]
"""
from __future__ import annotations

import argparse
import copy
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import docx  # noqa: E402
from docx.table import Table  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

from xtracts.config import PROJECT_ROOT, load_config, resolve  # noqa: E402

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
FIELDS = ["MERCHANT", "DATE", "TOTAL", "TAX"]
NICE = {"MERCHANT": "Merchant", "DATE": "Date", "TOTAL": "Total", "TAX": "Tax"}


# ------------------------------------------------------------------ docx helpers
def set_runs(p, texts):
    runs = p.runs
    if not runs:
        p.add_run("")
        runs = p.runs
    for i, r in enumerate(runs):
        r.text = texts[i] if i < len(texts) else ""
    for t in texts[len(runs):]:
        runs[-1].text += t


def set_body(p, body):
    runs = p.runs
    if len(runs) > 1 and runs[0].bold:
        set_runs(p, [runs[0].text, " " + body.lstrip()])
    else:
        set_runs(p, [body])


def set_cell(cell, text):
    set_runs(cell.paragraphs[0], [str(text)])
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def find_par(doc, startswith):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise KeyError(startswith)


def find_table(doc, first_cell):
    for t in doc.tables:
        if t.rows and t.rows[0].cells[0].text.strip().startswith(first_cell):
            return t
    raise KeyError(first_cell)


class Inserter:
    """Adds paragraphs / tables one after another, starting after an anchor element."""

    def __init__(self, doc, anchor_elm, protos):
        self.doc, self.cur, self.protos = doc, anchor_elm, protos

    def par(self, kind, texts):
        new = copy.deepcopy(self.protos[kind]._p)
        self.cur.addnext(new)
        self.cur = new
        p = Paragraph(new, self.protos[kind]._parent)
        set_runs(p, texts if isinstance(texts, list) else [texts])
        return p

    def blank(self):
        new = copy.deepcopy(self.protos["blank"]._p)
        self.cur.addnext(new)
        self.cur = new

    def table(self, caption_key, title, header, rows, ratios):
        self.par("caption", [f"Table @{caption_key}: ", title])
        proto = self.protos["table"]
        tbl = copy.deepcopy(proto._tbl)
        self.cur.addnext(tbl)
        self.cur = tbl
        trs = tbl.findall(W + "tr")
        hdr, body = copy.deepcopy(trs[0]), copy.deepcopy(trs[1])
        for tr in trs:
            tbl.remove(tr)
        n = len(header)
        total = sum(int(g.get(W + "w")) for g in tbl.find(W + "tblGrid"))
        widths = [int(total * r / sum(ratios)) for r in ratios]

        def fit(tr):
            tcs = tr.findall(W + "tc")
            while len(tcs) < n:
                tr.append(copy.deepcopy(tcs[-1]))
                tcs = tr.findall(W + "tc")
            for extra in tcs[n:]:
                tr.remove(extra)
            for tc, w in zip(tr.findall(W + "tc"), widths):
                tcw = tc.find(W + "tcPr").find(W + "tcW")
                tcw.set(W + "w", str(w))
                tcw.set(W + "type", "dxa")

        grid = tbl.find(W + "tblGrid")
        for g in list(grid):
            grid.remove(g)
        for w in widths:
            g = copy.deepcopy(proto._tbl.find(W + "tblGrid")[0])
            g.set(W + "w", str(w))
            grid.append(g)
        for tr_src, vals in [(hdr, header)] + [(body, r) for r in rows]:
            tr = copy.deepcopy(tr_src)
            fit(tr)
            tbl.append(tr)
        t = Table(tbl, proto._parent)
        for ri, vals in enumerate([header] + rows):
            for ci, v in enumerate(vals):
                set_cell(t.rows[ri].cells[ci], v)
        self.blank()
        return t


# ------------------------------------------------------------------ numbers
def pct(x):
    return f"{100 * x:.1f}%"


def counts(entity_level: dict) -> dict:
    """TP/FP/FN per field recovered from precision, recall and support, and checked against the micro scores."""
    out = {}
    for f in FIELDS:
        m = entity_level["fields"][f]
        tp = round(m["recall"] * m["support"])
        fn = m["support"] - tp
        fp = round(tp / m["precision"]) - tp if m["precision"] else 0
        out[f] = (tp, fp, fn)
    TP, FP, FN = (sum(out[f][k] for f in FIELDS) for k in range(3))
    mic = entity_level["micro"]
    ok = abs(TP / max(1, TP + FP) - mic["precision"]) < 0.002 and abs(TP / max(1, TP + FN) - mic["recall"]) < 0.002
    return {"fields": out, "micro": (TP, FP, FN), "consistent": ok}


def prf(tp, fp, fn):
    p = tp / (tp + fp) if tp + fp else 0.0
    r = tp / (tp + fn) if tp + fn else 0.0
    return p, r, (2 * p * r / (p + r) if p + r else 0.0)


# ------------------------------------------------------------------ the refinement
def build(src: Path, out: Path, metrics: dict | None, template: Path | None) -> dict:
    doc = docx.Document(str(src))
    cfg = load_config()
    hp = cfg["training"]
    full = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)

    # ---- facts that already stand in the report (nothing is typed in from memory)
    m = re.search(r"best epoch was (\d+) \(validation micro-F1 ([\d.]+)\); training took (\d+) seconds", full)
    best_epoch, val_f1, secs = (m.group(1), m.group(2), m.group(3)) if m else ("[epoch]", "[val F1]", "[seconds]")
    m = re.search(r"(\d+) / (\d+) / (\d+) \(human-reviewed receipts only; (\d+) reviewed", full)
    n_tr, n_va, n_te, n_rev = m.groups() if m else ("[n]",) * 4
    cmp_t = find_table(doc, "Model")  # first 'Model' table is Table 7; pick the comparison table below
    for t in doc.tables:
        if t.rows[0].cells[0].text.strip() == "Model" and t.rows[0].cells[1].text.startswith("Value exact"):
            cmp_t = t
    sysrows = {r.cells[0].text.strip(): [c.text.strip() for c in r.cells] for r in cmp_t.rows[1:]}
    rules, llm = sysrows.get("Rules"), sysrows.get("LayoutLMv3")

    protos = {
        "heading": find_par(doc, "7.4 Model Selection"),
        "body": find_par(doc, "The models, engines and methods"),
        "lead": find_par(doc, "Field-level extraction quality."),
        "eq": find_par(doc, "F1 = 2") if False else [p for p in doc.paragraphs if p.text.strip().startswith("F1 = 2")][0],
        "caption": find_par(doc, "Table 8:"),
        "step": find_par(doc, "1. Data Collection:"),
        "table": find_table(doc, "Model"),
        "blank": None,
    }
    blank = docx.oxml.OxmlElement("w:p")
    protos["blank"] = Paragraph(blank, protos["body"]._parent)

    changed: list[str] = []          # text of every paragraph written or rewritten here (for the similarity check)

    def rewrite(start, text):
        p = find_par(doc, start)
        set_body(p, text) if (p.runs and p.runs[0].bold and len(p.runs) > 1) else set_runs(p, [text])
        changed.append(text)

    # ---- 0. fix the dataset table (rows were one position off)
    t2 = find_table(doc, "Parameter")
    while t2.rows[1].cells[0].text.strip() != "Dataset name":
        t2 = next(t for t in doc.tables if t is not t2 and t.rows[0].cells[0].text.strip() == "Parameter"
                  and t.rows[1].cells[0].text.strip() == "Dataset name")
    rows2 = [
        ("Dataset name", "SROIE-style Malaysian receipts (SROIE: Scanned Receipts OCR and Information Extraction, ICDAR 2019)"),
        ("Source", "Receipt image set provided for this project [state where your copy was downloaded from]; the competition is described in [12]"),
        ("Data modality", "Receipt images; OCR words with boxes (our own PaddleOCR run); field labels made by rule-based suggestion plus human review"),
        ("Number of records (images)", "2,586 image files; 865 unique receipts after cleaning and de-duplication (1,721 removed: 117 exact copies and 1,604 near-duplicates)"),
        ("Training / validation / testing receipts", f"{n_tr} / {n_va} / {n_te} (human-reviewed receipts only; {n_rev} reviewed in total)"),
        ("Number of features", "Per word: text, bounding box (normalised to 0-1000) and, through the page image, visual content"),
        ("Target variable(s)", "Merchant, date, total, tax (BIO-tagged word spans)"),
        ("Missing values", "Unreadable or under-300-pixel images were excluded; receipts the reviewer marked unusable were excluded from training"),
        ("Duplicate records", "117 exact and 1,604 near-duplicate images found and removed, keeping one image per cluster (largest cluster: 54 images)"),
        ("Licence / terms of use", "[Check and state the dataset's terms of use; the SROIE terms are on the competition site cited in [12]]"),
    ]
    for i, (a, b) in enumerate(rows2, start=1):
        set_cell(t2.rows[i].cells[0], a)
        set_cell(t2.rows[i].cells[1], b)

    # ---- 1. rewrite generic paragraphs in new wording (the sample report's wording is shared by every student)
    rewrite("Receipts and invoices are among", "A receipt is a small document with a large role: it is the proof of almost every everyday purchase, from groceries and meals to medicines and repairs. It names the seller, gives the date and shows what was paid. Although more payments are made digitally each year, many of these proofs are still paper slips, or phone photographs and scans of them. In that form the information is locked inside pixels: a program cannot add it up, search it or chart it until the characters have been read and organised into fields.")
    rewrite("Converting receipts into structured data", "Turning receipts into data is useful well beyond the individual. People tracking a household budget, employees claiming reimbursements, small shops preparing accounts and auditors checking claims all depend on the right merchant, date and amount being recorded. Typing these values by hand costs time, invites mistakes and scales badly as the pile of receipts grows. A system that reads the image and proposes the key fields can reduce that effort and make the resulting records more dependable.")
    rewrite("Automating this task is difficult", "The task is harder than it first looks. No two merchants print the same layout, and fonts, languages, logos and the position of the total all vary. Thermal paper fades, and photographs may be tilted, blurred, unevenly lit or small. Whatever the text-recognition step gets wrong is passed on to every later step. A receipt also holds many numbers and dates that are not the one wanted (item prices, subtotals, tax lines, the cash handed over and the change), so choosing the amount that was actually due depends on context and not on how a number looks. For these reasons receipt understanding is studied as a research task of its own and has public benchmarks [1].")
    rewrite("Artificial intelligence and machine learning contribute", "Machine learning helps in two places. Optical character recognition (OCR) first finds the text in the picture and reads it; older engines [2] and deep text detectors and recognisers [3], [4] are widely used. Information extraction then decides which of the recognised words carry meaning, normally using both the words and where they sit on the page. Methods range from hand-written rules and regular expressions, through supervised classifiers, to large pre-trained language and layout models [5], [6].")
    rewrite("Real-world problem.", "Anyone who spends money and later has to account for it (a student sharing costs, a freelancer claiming expenses, a shopkeeper doing the books) ends up with a pile of receipts. They are easy to lose, hard to sort and tedious to enter into a spreadsheet, so records are often late or incomplete, which makes budgeting, tax filing and reimbursement harder.")
    rewrite("Limitations of existing approaches.", "Manual entry is slow and error-prone. Plain OCR tools return an undifferentiated block of text and leave the user to work out which line is the seller, the date or the amount. Template-based and rule-based extractors are effective on the layouts they were written for but must be rewritten whenever a new format appears. Commercial cloud receipt-parsing services work well, but they are charged per use, need financial documents to be uploaded to a third party and cannot be inspected or adapted.")
    rewrite("Need for an automated, intelligent solution.", "Reading a receipt is partly a vision problem and partly a language problem, and the right answer depends on position as much as on content. A system that combines OCR with a learned model of text and layout can cope with layouts it was not explicitly programmed for, while a confirmation step lets the user correct mistakes before they enter the records.")
    rewrite("Expected benefits.", "The system is expected to (i) reduce typing, (ii) lower transcription mistakes, (iii) put expenses into a structured form that can be queried and summarised, and (iv) give students a complete, reproducible example of an applied document-AI pipeline, from raw images to a working application.")
    rewrite("Research relevant to Xtracts", "The research relevant to Xtracts falls into four groups: reading text from images, detecting where text lies, benchmark datasets for receipts and forms, and models that extract structured fields from recognised text. Table 1 lists representative works with the data they used, their technique, main findings and limitations.")
    rewrite("Discussion. Four families", "Four approaches stand out. First, the classical pipeline reads the text with OCR and then applies rules, patterns or classifiers; it is easy to understand but is thrown off by OCR errors and by layouts it has not seen. Second, deep OCR splits the job into detecting text [4] and recognising it [3], which copes better with varied fonts and image quality. Third, layout-aware transformers [6], [7], [11] take each word together with its position, and often the page image, so they can learn cues such as ‘the number beside the word TOTAL’. Fourth, OCR-free models [8] map the image straight to structured output and so avoid passing OCR mistakes forward, at the price of heavy computation.")
    rewrite("The main advantages of the modular", "The modular approach is transparent and light on resources, but it passes errors along and can be brittle. Layout-aware and OCR-free transformers give better extraction in the cited studies, yet they need fine-tuning on labelled data with considerable compute and are harder to inspect when they fail.")
    rewrite("Identified gap.", "Most published work aims to raise benchmark accuracy for the extraction step. Less attention goes to a small, reproducible pipeline that carries the extracted fields into an application ordinary users can operate, and that also reports how well each stage performs on hardware available to students.")
    rewrite("Feature description.", "For an image-and-text problem the inputs are not columns of a table but the receipt picture, the recognised words and the box around each word. Table 3 lists every input and every target with its type and role.")
    rewrite("Version control records", "Version control keeps a dated record of every change to a project, so that any change can be inspected, compared or undone and several people can work side by side [17]. Xtracts was kept in a Git repository holding the source code, a README with run instructions and a .gitignore that leaves out data, models and virtual environments. Work was committed in small steps, one per completed stage: dataset preparation, OCR, annotation tools, training code, evaluation, application and report.")
    rewrite("Experiment tracking means", "Experiment tracking means keeping a record of every run (its data, settings, metrics and output files) so that results can be repeated and compared later [18]. Tools such as MLflow or Weights & Biases are commonly used. MLflow was chosen for Xtracts because it runs locally without an account, stores its records in one SQLite file and needs only a few lines of code.")
    rewrite("AutoML and MLOps are optional", "AutoML and MLOps are optional topics for a mini-project. This section states how far each applies to Xtracts; whatever was not done is described as future work.")
    rewrite("AutoML automates", "AutoML tools search automatically over algorithms and hyper-parameters [19]. None was used here: the OCR engine is pre-trained, and a single LayoutLMv3 configuration with fixed hyper-parameters was fine-tuned. A search with a tool such as Optuna is left as future work.")
    rewrite("MLOps applies", "MLOps brings software-engineering practice (versioned code and data, tracked experiments, automated tests, deployment pipelines and monitoring) to the machine-learning life cycle [20]. Implemented in Xtracts: Git, a pinned requirements file, one YAML configuration file, pytest tests covering data processing, the rule-based labeler, the training loop, evaluation, the database and the application pages, and MLflow tracking of evaluation runs. Not implemented: a model registry, CI/CD, containers and monitoring after deployment.")
    rewrite("To study the receipt OCR", "To review how receipt OCR and key-information extraction are done today, including the SROIE benchmark.")
    rewrite("To build a pipeline that converts", "To build a pipeline that reads a receipt image, recognises its text and converts it into structured merchant, date, total and tax fields.")
    rewrite("To document the limitations", "To state the limitations of the approach honestly and to suggest realistic improvements.")
    rewrite("Security and privacy.", "Receipts reveal spending habits, so later versions could encrypt stored images, add user accounts and run the whole pipeline on the user's own device.")
    rewrite("Line-item and category extraction.", "Reading each purchased item with its price, and assigning categories such as food or transport, would let spending be analysed by type and not only by merchant.")
    rewrite("Improved image handling.", "Straightening and enhancing phone photographs, for example with perspective correction and contrast improvement, could help on images that are skewed or poorly lit.")
    rewrite("Broader data.", "More verified receipts, receipts from other collections such as CORD [10] and locally collected receipts would show whether the model generalises beyond these merchants.")
    rewrite("Automated retraining and monitoring.", "Corrections that users make in the application could be collected to retrain the model, with monitoring of accuracy over time.")
    p = find_par(doc, "Exploratory analysis concentrated")
    set_runs(p, [p.text.replace(" [Insert one or two EDA figures here, e.g. words per receipt, if you produce them.]", "")])

    # ---- 2. design choices: reasons and alternatives (after the models table)
    models_tbl = [t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Model" and t.rows[0].cells[1].text.strip() == "Purpose"][0]
    ins = Inserter(doc, models_tbl._tbl, protos)
    ins.blank()
    intro = ("Why these choices, and what else could have been used. Each main decision is summarised in Table @design with the reason it "
             "was taken and the alternatives that exist. The alternatives were not benchmarked in this project: the choices rest on the "
             "requirements of a small, offline, CPU-and-free-GPU project, and the comparisons in the last column are reasons, not measurements.")
    ins.par("lead", ["Why these choices, and what else could have been used. ", intro.split(". ", 1)[1]])
    changed.append(intro)
    design = [
        ("Duplicate detection", "SHA-256 plus perceptual hashes (pHash, dHash), Hamming distance ≤ 6",
         "Catches exact copies and near-copies (re-scans, resized or re-compressed images) on a CPU with no training; the threshold could be checked by eye on contact sheets.",
         "Image embeddings from a CNN or CLIP-style model; local feature matching (ORB, SIFT); comparing file names or metadata.",
         "Embeddings and feature matching are heavier and need their own tuning; file names and metadata are unreliable for scanned files."),
        ("OCR engine", "PaddleOCR 2.9.1 (pre-trained) [22]",
         "Open source, runs offline on a CPU and returns text, a polygon box and a confidence for each line; a larger detection limit keeps small print on tall receipts readable.",
         "Tesseract [2]; EasyOCR; TrOCR (recognition only, needs a separate detector); cloud OCR services such as Google Cloud Vision, Amazon Textract or Azure Document Intelligence.",
         "Cloud services need internet, per-use fees and uploading financial documents; Tesseract and EasyOCR were not tried, so no accuracy claim is made against them."),
        ("Field extractor", "LayoutLMv3-base fine-tuned for token classification [11]",
         "Reads the words, their boxes and the page image together, so position cues such as ‘amount beside TOTAL’ can be learned; the base model fine-tunes on a few hundred receipts and trained in minutes on a free GPU; it uses patch embeddings, so no separate image backbone is needed.",
         "LayoutLM [6] and LayoutLMv2 [7]; a text-only BERT-style tagger [5]; Donut, which needs no OCR [8]; graph-based models [9]; LayoutLMv3-large.",
         "A text-only tagger ignores layout; Donut is generative, heavier and gives no word-level labels to review; larger models need more data and memory than the project had."),
        ("Baseline", "Rule-based extractor (regular expressions, keywords, row grouping)",
         "Needs no training data, can be read and corrected line by line, and doubles as the source of label suggestions; it gives a yardstick against which the learned model is judged.",
         "Per-merchant templates; classical machine learning on hand-made features (CRF, gradient boosting); fuzzy matching.",
         "Templates do not generalise to new merchants; classical models on hand-made features were not tried because of time."),
        ("Ground truth", "Rule suggestions corrected by a human in a Streamlit review app",
         "The images had no field annotations, and labelling 865 receipts from scratch is slow; correcting a suggestion is much quicker than typing a label.",
         "Fully manual labelling in a general tool (for example Label Studio); the official SROIE annotations; crowd-sourcing; labelling with a large language model.",
         "No official labels exist for this image set; general tools need setup and do not know the word-box format used here. The cost is a possible bias towards the suggestion (see limitations)."),
        ("Data split", f"Merchant-grouped split (GroupShuffleSplit, seed {cfg['seed']}): {n_tr} / {n_va} / {n_te} receipts",
         "Receipts of one merchant look alike; keeping each merchant in one set stops the model from scoring well by remembering layouts.",
         "Random split; k-fold cross-validation; split by time.",
         "A random split would inflate scores; k-fold would multiply the training time on a small GPU budget."),
        ("Long receipts", f"Windows of up to {hp['max_length']} tokens with {hp['overlap_words']}-word overlap",
         "The total and tax sit at the bottom, which plain truncation would cut off; overlap keeps a field whole in at least one window.",
         "Truncation; shrinking the text; long-context document models.",
         "Truncation removes exactly the fields that matter."),
        ("Training recipe", f"Custom loop: AdamW, lr {hp['learning_rate']}, linear warm-up and decay, batch {hp['batch_size']} × {hp['grad_accum_steps']} accumulation, mixed precision, early stopping on validation micro-F1",
         "A standard fine-tuning recipe; accumulation and mixed precision fit a free GPU, and early stopping limits over-fitting on a small data set; a custom loop gave control over window merging and the entity-level validation score.",
         "The Hugging Face Trainer; parameter-efficient fine-tuning (LoRA); a hyper-parameter search (Optuna); a class-weighted loss.",
         "Not tried for lack of time; the settings were fixed before training and not tuned on the test set."),
        ("Evaluation metrics", "Entity-level precision, recall, F1 and value exact match",
         "Most words belong to no field, so plain accuracy would look high even for a useless model; entity-level scores follow the usual tagging convention and value exact match shows what the user would see.",
         "Token-level accuracy or F1; CER and WER; fuzzy or edit-distance matching; the seqeval library.",
         "CER/WER need transcriptions that do not exist here; seqeval would not install in our environment, so the same definition was implemented and unit-tested."),
        ("Experiment tracking", "MLflow with a SQLite backend [18]",
         "Local, no account, one file, a browser UI for comparing runs.",
         "Weights & Biases; TensorBoard; a CSV or spreadsheet log.",
         "Weights & Biases needs an online account; TensorBoard is better for loss curves than for comparing evaluation tables."),
        ("Application", "Streamlit",
         "Pure Python, so upload, editable fields and charts can be built quickly in the same language as the model code.",
         "Flask or FastAPI with a JavaScript front end; Django; Gradio.",
         "A separate front end is more work than a mini-project needs; Gradio gives less control over a multi-page dashboard."),
        ("Database", "SQLite",
         "A single file, built into Python, no server to run; enough for one user.",
         "PostgreSQL or MySQL; a CSV or Excel file; a cloud database.",
         "Server databases matter once several users share data; a CSV file cannot enforce constraints such as duplicate protection."),
    ]
    ins.table("design", "Design choices, reasons and alternatives",
              ["Component", "Chosen", "Why it was chosen", "Alternatives available", "Why not used here"],
              [list(r) for r in design], [1.55, 1.7, 2.7, 2.4, 2.3])
    changed += [r[2] + " " + r[4] for r in design]

    # ---- 3. evaluation: formulas + procedure + counts + normalisation (inside 7.6)
    eq4 = protos["eq"]
    ins = Inserter(doc, eq4._p, protos)
    ins.par("lead", ["Pooling over fields (micro-average). ", "Besides one score per field, a single overall score is reported. The counts of "
                     "all four fields are added together before dividing, so that every entity counts once:"])
    ins.par("eq", ["\t", "Micro-Precision = Σ TP / Σ (TP + FP)", "\t(5)"])
    ins.par("eq", ["\t", "Micro-Recall = Σ TP / Σ (TP + FN)", "\t(6)"])
    ins.par("lead", ["Value exact match. ", "This is a second, user-oriented measure. For each receipt and field the predicted text and the "
                     "human text are first turned into a comparable value (rules below) and then compared. A receipt with no human value "
                     "and no prediction for a field counts as correct."])
    ins.par("eq", ["\t", "Value exact match = receipts whose normalised values agree / receipts evaluated", "\t(7)"])
    changed.append("pooling over fields micro-average value exact match")

    metrics_tbl = find_table(doc, "Metric")
    ins = Inserter(doc, metrics_tbl._tbl, protos)
    ins.blank()
    ins.par("lead", ["Evaluation procedure, step by step. ", f"The scores in Section 12 were produced by scripts/07_evaluate.py in the following steps."])
    steps = [
        ("1. Test set: ", f"The {n_te} test receipts (merchants not present in training or validation) are read from test.jsonl. Each holds the OCR words, their boxes scaled to 0-1000 and the human BIO labels, which serve as the gold standard."),
        ("2. LayoutLMv3 prediction: ", f"Each receipt is cut into windows of at most {hp['max_length']} tokens with {hp['overlap_words']} words of overlap. For every word the model gives a probability for each of the nine tags (O and B-/I- for the four fields); the first sub-word token of the word decides, and the most probable tag is taken. If a word lies in several windows, the prediction from the window in which it is farthest from the edge is kept."),
        ("3. Rule-based prediction: ", "The same OCR words go through the current rule-based labeler, which returns tags in the same format. Both systems are therefore scored on identical input and identical gold labels."),
        ("4. From tags to entities: ", "A B-X tag opens an entity of field X, following I-X tags extend it, and a stray I-X opens one. An entity is the triple (field, first word, last word)."),
        ("5. Matching: ", "A predicted entity is a true positive (TP) only if an identical triple exists in the gold labels. A predicted entity without a gold match is a false positive (FP); a gold entity that was not predicted exactly is a false negative (FN). A prediction that covers the right value but with one word too many or too few is therefore both an FP and an FN."),
        ("6. Scores: ", "TP, FP and FN are summed over all test receipts, separately for each field, and precision, recall and F1 follow from equations (2)-(4). The overall score pools the four fields with equations (5) and (6) and then uses equation (4)."),
        ("7. Value exact match: ", "Per receipt and field the entity text is normalised and compared with the gold value (next paragraph) to give equation (7); the figure in Section 12 is the mean over the four fields."),
        ("8. Outputs: ", "metrics.json (all numbers), field_metrics.csv (one row per system and field), per_receipt.csv (gold and predicted value for every receipt and field, used for error analysis), summary.md, and one MLflow run with the parameters, metrics and these files."),
    ]
    for lead, rest in steps:
        ins.par("step", [lead, rest])
        changed.append(rest)
    ins.par("lead", ["Normalisation used for value exact match. ", "Merchant: lower-case, keeping letters and digits only, so ‘FRESH MART SDN. BHD.’ and "
                     "‘Fresh Mart Sdn Bhd’ agree. Date: day, month and year are parsed (day-first, then month-first if the first reading is not a "
                     "valid date) and written as YYYY-MM-DD. Total and tax: currency symbols are removed and the number is rounded to two decimals, so "
                     "‘RM12.50’ and ‘12.5’ agree. If a value cannot be parsed, the text with spaces removed is compared instead."])

    # counts table + worked example
    systems_counts = {}
    if metrics:
        for name, s in metrics["systems"].items():
            c = counts(s["entity_level"])
            systems_counts[name] = c
    rows_c = []
    for name, c in systems_counts.items():
        if not c["consistent"]:
            continue
        for f in FIELDS + ["ALL"]:
            tp, fp, fn = c["micro"] if f == "ALL" else c["fields"][f]
            p, r, f1 = prf(tp, fp, fn)
            rows_c.append([name, "All fields (micro)" if f == "ALL" else NICE[f], str(tp), str(fp), str(fn), pct(p), pct(r), pct(f1)])
    ins.blank()
    if rows_c:
        ins.table("counts", f"True positives, false positives, false negatives and scores on the {n_te} test receipts",
                  ["System", "Field", "TP", "FP", "FN", "Precision", "Recall", "F1"], rows_c, [1.4, 1.8, 0.7, 0.7, 0.7, 1.2, 1.2, 1.2])
        name0 = next(n for n, c in systems_counts.items() if c["consistent"])
        tp, fp, fn = systems_counts[name0]["fields"]["TOTAL"]
        p, r, f1 = prf(tp, fp, fn)
        ex = (f"Worked example. For the Total field of the {name0} system there are {tp} true positives, {fp} false positive(s) and {fn} false negatives. "
              f"Precision = {tp} / ({tp} + {fp}) = {pct(p)}; recall = {tp} / ({tp} + {fn}) = {pct(r)}; F1 = 2 × {p:.3f} × {r:.3f} / ({p:.3f} + {r:.3f}) = {pct(f1)}. "
              "The overall score pools all four fields in the same way (last rows of Table @counts).")
        ins.par("lead", [ex.split(". ", 1)[0] + ". ", ex.split(". ", 1)[1]])
        changed.append(ex)
    else:
        ins.par("body", ["[Counts table: run scripts/07_evaluate.py and then this script again with reports/evaluation/metrics.json in place]"])

    # ---- 4. stage-wise execution and outputs (new 7.7, before section 8)
    ins.par("heading", ["7.7 Stage-wise Execution and Outputs"])
    ins.par("body", ["Table @stages follows the work in the order it was done, with what was carried out at each stage, what it produced and the measured result. "
                     "Figures are taken from the runs of this project."])
    if rules and llm:
        res = f"Test split: rule baseline micro-F1 {rules[4]}, LayoutLMv3 {llm[4]}; mean value exact match {rules[1]} and {llm[1]}"
    else:
        res = "[test results: run scripts/07_evaluate.py]"
    stages = [
        ["1. Cleaning and de-duplication", "Checked every file; computed SHA-256 and perceptual hashes; grouped copies around one kept image per cluster; inspected contact sheets. An earlier chaining method had merged unrelated receipts into clusters of up to 222 images and was replaced.",
         "Clean images, manifest, duplicate report, contact sheets", "2,586 files → 865 unique; 117 exact and 1,604 near-duplicates removed; largest cluster 54 images"],
        ["2. OCR", "Ran PaddleOCR on every clean image; sorted lines into reading order; split lines into words with proportional boxes; saved confidence.",
         "One JSON file per receipt (words, boxes, confidence)", "865 OCR files; mean line confidence 0.949 in the OCR summary file"],
        ["3. Ground truth", "Rule-based labeler proposed merchant, date, total and tax; a reviewer corrected them in the review app. The date and tax rules were improved in several rounds after studying the corrections.",
         "One label file per receipt (suggested labels, corrected labels, flags)", f"{n_rev} receipts reviewed; date suggestions found rose from 31% to about 92% and tax from 38% to about 83% of unreviewed receipts (found, not verified accurate)"],
        ["4. Dataset build", "Converted reviewed labels into words, 0-1000 boxes and BIO tags; split by merchant.",
         "train.jsonl, val.jsonl, test.jsonl", f"{n_tr} / {n_va} / {n_te} receipts"],
        ["5. Training", f"Fine-tuned LayoutLMv3-base on a GPU; validated after each epoch; kept the best checkpoint.",
         "Model folder, training log, validation metrics, summary.json", f"Best epoch {best_epoch}; validation micro-F1 {val_f1}; training time {secs} s"],
        ["6. Evaluation", "Scored both systems on the test receipts with the procedure of Section 7.6.",
         "metrics.json, field_metrics.csv, per_receipt.csv, summary.md, MLflow run", res],
        ["7. Application", "Upload → OCR → LayoutLMv3 → editable fields → SQLite → dashboard and history; checks for invalid files and duplicates.",
         "expenses.db, copies of saved images", "Working local Streamlit application covered by automated page tests"],
    ]
    ins.table("stages", "Stages, activities, outputs and results", ["Stage", "What was done", "Output", "Measured result"], stages, [1.3, 3.6, 2.2, 2.6])
    changed += [s[1] for s in stages]

    # ---- 5. results discussion: honest comparison
    rules_better = bool(rules and llm) and float(rules[4].rstrip("%")) > float(llm[4].rstrip("%"))
    if rules and llm and rules_better:
        why = (f"On the {n_te} test receipts the rule-based extractor scored a micro-F1 of {rules[4]} against {llm[4]} for LayoutLMv3, "
               f"while the mean value exact match was {rules[1]} against {llm[1]}. Several circumstances, none of them tested separately, may explain this. "
               f"(i) Only {n_tr} receipts were available for training, which is little for a model with many parameters. "
               "(ii) The human labels began as the rules' own suggestions, so the model was taught what the rules produce and cannot easily exceed them. "
               "(iii) The rules were refined while studying reviewed receipts, some of which fall in the test set, so their score is optimistic. "
               "(iv) Entity-level F1 demands identical word spans, whereas value exact match forgives small differences such as a currency symbol; the smaller gap in the second measure suggests that part of the model's loss comes from span boundaries and not from wrong values. "
               "(v) With 54 receipts, one receipt equals about 1.9 percentage points per field, so differences of a few points are not conclusive.")
        lead = "Why the rule baseline was not beaten. "
    elif rules and llm:
        why = (f"On the {n_te} test receipts LayoutLMv3 scored a micro-F1 of {llm[4]} against {rules[4]} for the rule-based extractor (mean value exact match {llm[1]} against {rules[1]}). "
               "A plausible reason, not tested separately, is that the model learns layout cues from examples instead of relying on fixed patterns. Three cautions apply: "
               "(i) with 54 receipts one receipt equals about 1.9 percentage points per field, so small gaps are not conclusive; "
               "(ii) the human labels began as the rules' suggestions, which may have anchored the annotator; "
               "(iii) the rules were refined while studying reviewed receipts, some of which fall in the test set, which favours the rules, so the true margin may be larger.")
        lead = "Why LayoutLMv3 scored higher. "
    if rules and llm:
        ins2 = Inserter(doc, find_par(doc, "Best-performing method")._p, protos)
        ins2.par("lead", [lead, why])
        ins2.par("lead", ["Where LayoutLMv3 is still preferable. ", "The rules were written and tuned for the receipt styles seen in this collection and need new rules for every new layout; LayoutLMv3 learns from labelled examples, can be improved by adding more corrected receipts and can learn merchant-specific cues without code changes. These advantages are plausible but were not measured here."])
        changed += [why, "where LayoutLMv3 is preferable"]
    rewrite("Limitations.", "Only four fields are extracted; line items and categories are not. Ground truth came from a single reviewer correcting rule suggestions, and the first receipts reviewed were chosen because they were flagged. The rules were refined while studying reviewed receipts, some of which lie in the test set, which favours the rule baseline. The data are Malaysian receipts, so other layouts and languages are untested. The verified set is a few hundred receipts and the test set 54, so the uncertainty is large. OCR quality was not measured separately, and the word boxes inside a line are estimated proportionally.")
    for p in doc.paragraphs:
        if p.text.startswith("The final system turns receipt images") and rules and llm:
            verdict = ("LayoutLMv3 therefore did not surpass the rules in entity-level F1 on this small test set, although the two are close on value-level agreement."
                       if rules_better else "LayoutLMv3 therefore scored higher than the rules in entity-level F1 on this small test set.")
            set_runs(p, [f"The final system turns receipt images into structured records and expense analytics. On the {n_te} held-out, human-verified test receipts the rule-based baseline reached a micro-F1 of {rules[4]} and a mean value exact match of {rules[1]}, and the fine-tuned LayoutLMv3 model reached {llm[4]} and {llm[1]}. {verdict} The main limitations are the small verified data set, a single annotator, rules tuned on receipts that include test receipts, the lack of a separate OCR evaluation and the restriction to four fields. Overall, the work demonstrates the practical value of automated receipt understanding for expense management, and the future-scope items in Section 13 outline directions for improvement."])
            changed.append(p.text)

    if metrics and len(metrics["systems"]) > 1:
        field_tbl = [t for t in doc.tables if t.rows[0].cells[0].text.strip().startswith("Field (")][0]
        rows_f = []
        for name, sy in metrics["systems"].items():
            for f in FIELDS:
                e, v = sy["entity_level"]["fields"][f], sy["value_level"][f]
                rows_f.append([name, NICE[f], pct(e["precision"]), pct(e["recall"]), pct(e["f1"]), pct(v["value_exact"])])
        ins3 = Inserter(doc, field_tbl._tbl, protos)
        ins3.blank()
        ins3.table("fieldall", "Field-level results of both systems on the test receipts",
                   ["System", "Field", "Precision", "Recall", "F1", "Value exact match"], rows_f, [1.6, 1.4, 1.3, 1.3, 1.3, 1.8])

    # ---- 6. renumber tables in order of appearance and resolve @tokens
    caps = []
    for el in doc.element.body.iterchildren(W + "p"):
        text = "".join(t.text or "" for t in el.iter(W + "t"))
        mm = re.match(r"Table (@\w+|\d+):", text)
        if mm:
            caps.append(mm.group(1))
    newnum = {c: str(i) for i, c in enumerate(caps, start=1)}

    def fix(text):
        return re.sub(r"Table (@\w+|\d+)", lambda m_: "Table " + newnum.get(m_.group(1), m_.group(1)), text)

    def fix_runs(par):
        for r in par.runs:
            if "Table" in r.text:
                r.text = fix(r.text)

    for p in doc.paragraphs:
        fix_runs(p)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    fix_runs(p)
    def drop_highlight(par):
        """Remove the template's yellow ‘to be filled’ marking from text that has been filled in."""
        for r in par.runs:
            if re.search(r"\[(?!\d)", r.text):          # a real placeholder such as [state ...]: keep it marked
                continue
            if r.font.highlight_color is not None:
                r.font.highlight_color = None
            rpr = r._r.rPr
            if rpr is not None:
                for shd in rpr.findall(W + "shd"):
                    if (shd.get(W + "fill") or "").upper() in ("FFFF00", "FFFF66", "FFFF99"):
                        rpr.remove(shd)
    for p in doc.paragraphs:
        drop_highlight(p)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    drop_highlight(p)
    leftover = [p.text[:60] for p in doc.paragraphs if re.search(r"Table @", p.text)]
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))

    # ---- 7. similarity of the new/rewritten text with the sample report
    sim = None
    if template and template.exists():
        tdoc = docx.Document(str(template))
        ttext = " ".join(p.text for p in tdoc.paragraphs) + " " + " ".join(c.text for t in tdoc.tables for r in t.rows for c in r.cells)

        def sh(s, n=6):
            w = re.findall(r"[a-z0-9]+", s.lower())
            return {" ".join(w[i:i + n]) for i in range(max(0, len(w) - n + 1))}
        ts = sh(ttext)
        mine = set().union(*[sh(c) for c in changed])
        # paragraphs of the FINAL report that still share many phrases with the sample report
        shared = []
        for p in docx.Document(str(out)).paragraphs:
            ps = sh(p.text)
            if re.match(r"^\[\d+\]", p.text):
                continue                                   # references are meant to match
            if len(ps) >= 8 and len(ps & ts) / len(ps) >= 0.3:
                shared.append((len(ps & ts) / len(ps), p.text[:90]))
        sim = (len(mine & ts) / max(1, len(mine)), len(mine), sorted(shared, reverse=True))
    return {"leftover": leftover, "similarity": sim, "tables": len(caps), "counts": {k: v["consistent"] for k, v in systems_counts.items()},
            "have_results": bool(rules and llm)}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--in", dest="src", default="docs/Xtracts_Mini_Project_Report_UPDATED_NEW.docx")
    ap.add_argument("--out", default="docs/Xtracts_Mini_Project_Report_REFINED.docx")
    ap.add_argument("--metrics", default=None)
    ap.add_argument("--template", default="docs/Xtracts_Mini_Project_Report_TEMPLATE.docx")
    a = ap.parse_args()
    cfg = load_config()
    mpath = Path(a.metrics) if a.metrics else resolve(cfg["paths"]["eval_dir"]) / "metrics.json"
    metrics = json.loads(mpath.read_text()) if mpath.exists() else None
    res = build(resolve(a.src), resolve(a.out), metrics, resolve(a.template))
    print(f"Wrote {resolve(a.out)}  ({res['tables']} numbered tables)")
    print("  metrics.json used:", "yes" if metrics else "NO (counts table skipped)", "| counts consistent with micro scores:", res["counts"])
    if res["leftover"]:
        print("  unresolved table references:", res["leftover"])
    if res["similarity"]:
        print(f"  6-word phrases of the new/rewritten text that also occur in the sample report: {100 * res['similarity'][0]:.1f}% of {res['similarity'][1]}")
        if res["similarity"][2]:
            print("  paragraphs still close to the sample report's wording (consider rewording in your own words):")
            for frac, txt in res["similarity"][2][:15]:
                print(f"    {100 * frac:3.0f}%  {txt}")


if __name__ == "__main__":
    main()
