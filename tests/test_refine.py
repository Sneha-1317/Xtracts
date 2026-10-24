import importlib.util
from pathlib import Path

import pytest

pytest.importorskip("docx")
ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs" / "Xtracts_Mini_Project_Report_TEMPLATE.docx"


def _load(name):
    spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _entity(f1, sup):
    fields = {f: {"precision": 0.9, "recall": 0.9, "f1": f1, "support": sup, "exact_match": 0.9} for f in ("MERCHANT", "DATE", "TOTAL", "TAX")}
    return {"fields": fields, "micro": {"precision": 0.9, "recall": 0.9, "f1": f1}}


def test_counts_recover_tp_fp_fn_and_match_micro_scores():
    mod = _load("09_refine_report")
    ent = {"fields": {"MERCHANT": {"precision": 47 / 52, "recall": 47 / 54, "support": 54},
                      "DATE": {"precision": 1.0, "recall": 48 / 53, "support": 53},
                      "TOTAL": {"precision": 46 / 47, "recall": 46 / 54, "support": 54},
                      "TAX": {"precision": 41 / 44, "recall": 41 / 52, "support": 52}},
           "micro": {"precision": 182 / 191, "recall": 182 / 213}}
    c = mod.counts(ent)
    assert c["fields"]["TOTAL"] == (46, 1, 8) and c["micro"] == (182, 9, 31) and c["consistent"]


@pytest.mark.skipif(not TEMPLATE.exists(), reason="template not present")
def test_refine_adds_sections_without_breaking_numbering(tmp_path):
    from xtracts.config import load_config

    mk = _load("08_make_report")
    ref = _load("09_refine_report")
    cfg = load_config()
    g = {"splits": {"train": (319, {}), "val": (62, {}), "test": (54, {})}, "reviewed": 435, "summary": {"best_epoch": 14, "best_val_f1": 0.847, "train_seconds": 693},
         "per": None, "ocr_conf": None,
         "metrics": {"meta": {"split": "test", "n_receipts": 54, "generated": "now"},
                     "systems": {"Rules": {"entity_level": _entity(0.9, 54), "value_level": {f: {"value_exact": 0.85, "n": 54, "n_with_gold": 54} for f in ("MERCHANT", "DATE", "TOTAL", "TAX")}},
                                 "LayoutLMv3": {"entity_level": _entity(0.85, 54), "value_level": {f: {"value_exact": 0.86, "n": 54, "n_with_gold": 54} for f in ("MERCHANT", "DATE", "TOTAL", "TAX")}}}}}
    mk.build(TEMPLATE, tmp_path / "base.docx", g, cfg)
    res = ref.build(tmp_path / "base.docx", tmp_path / "refined.docx", g["metrics"], TEMPLATE)
    import docx
    import re
    d = docx.Document(str(tmp_path / "refined.docx"))
    caps = [int(m.group(1)) for p in d.paragraphs if (m := re.match(r"Table (\d+):", p.text))]
    assert caps == list(range(1, len(caps) + 1))                       # sequential, no duplicates
    assert not res["leftover"]
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Design choices, reasons and alternatives" in text and "Evaluation procedure, step by step" in text
    assert "7.7 Stage-wise Execution and Outputs" in text and "Why the rule baseline was not beaten" in text
